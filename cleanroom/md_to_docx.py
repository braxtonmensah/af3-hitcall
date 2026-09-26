"""Convert the preprint Markdown into a .docx that bioRxiv will accept.

bioRxiv takes PDF or Word. The manuscript is Markdown and pandoc is not installed on this machine, so
this does the conversion with python-docx, which is. It is deliberately narrow: it handles the subset of
Markdown the preprint actually uses (ATX headings, bold, italic, inline code, pipe tables, bullet and
numbered lists, horizontal rules, blockquotes) and does not try to be a general converter.

Two choices worth knowing about:
  - Inline `code` is rendered in a monospace run rather than dropped, because the manuscript uses it for
    filenames and identifiers that a reader needs to see exactly.
  - Pipe tables become real Word tables, because the preprint's results are in tables and a reviewer
    reading a wall of pipes is a reviewer who stops reading.

Usage:
    py -3.11 md_to_docx.py ../PREPRINT_DRAFT.md ../preprint_for_biorxiv.docx
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text):
    """Split a line into bold / italic / code / plain runs."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def is_table_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def main(src, dst):
    lines = open(src, encoding="utf-8").read().split("\n")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # pipe table: header, separator, then body rows
        if line.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            body = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                body.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for k, h in enumerate(header):
                cell = t.rows[0].cells[k]
                cell.text = ""
                add_runs(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in body:
                cells = t.add_row().cells
                for k in range(min(len(row), len(header))):
                    cells[k].text = ""
                    add_runs(cells[k].paragraphs[0], row[k])
            doc.add_paragraph()
            i = j
            continue

        if not line.strip():
            i += 1
            continue

        if line.strip() in ("---", "***", "___"):
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            text = re.sub(r"[*`]", "", m.group(2))
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            add_runs(p, line.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        m = re.match(r"^\s*[-*+]\s+(.*)", line)
        if m:
            add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue

        m = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if m:
            add_runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue

        # Join wrapped lines into one paragraph, which is what the Markdown means.
        buf = [line]
        j = i + 1
        while (j < len(lines) and lines[j].strip()
               and not lines[j].startswith(("#", "|", ">", "---"))
               and not re.match(r"^\s*([-*+]|\d+[.)])\s", lines[j])):
            buf.append(lines[j].rstrip())
            j += 1
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, " ".join(buf))
        i = j

    doc.save(dst)
    print("wrote", dst)
    print("%d paragraphs, %d tables" % (len(doc.paragraphs), len(doc.tables)))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: md_to_docx.py <in.md> <out.docx>")
    main(sys.argv[1], sys.argv[2])
